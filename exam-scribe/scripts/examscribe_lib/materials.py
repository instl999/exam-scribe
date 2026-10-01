"""Instructor materials: past papers, syllabus, lecture slides.

The script does the searching (splitting a paper into questions, ranking the
book sections each question most likely comes from, finding where the slides
use a key term); the model only makes small, checkable choices from a closed
list of candidates. Results feed the priority stars, and slide/book conflicts
are shown next to the notes ("follow your instructor on the exam").
"""
from __future__ import annotations

import math
import re
import secrets
import shutil
from collections import Counter
from pathlib import Path

from .common import ESError, Workspace, read_json, read_text, sha, split_sentences, write_json, write_text
from .esm import autofix, parse
from .ingest import section_paragraphs

STOP = set("""a an the and or but if then than that this these those is are was were be been being of to in on for
with by as at from into about over under between through during which what when where who whom why how all any
both each few more most other some such no nor not only own same so too very can will just should now also may might
must shall would could does did do has have had its it they them their there here you your we our he she his her
question questions answer explain describe calculate find show following given value values use using""".split())


def _tokens(text: str) -> list[str]:
    return [w for w in re.findall(r"[a-zÀ-ɏͰ-Ͽ一-鿿][a-z0-9À-ɏͰ-Ͽ一-鿿\-]{2,}",
                                  text.casefold()) if w not in STOP]


def _extract_text(path: Path) -> list[tuple[str, str]]:
    """[(page label, text)] for a PDF/text file."""
    suf = path.suffix.lower()
    if suf in (".pdf", ".epub", ".xps"):
        import pymupdf
        doc = pymupdf.open(str(path))
        return [(page.get_label() or str(i + 1), page.get_text("text", sort=True)) for i, page in enumerate(doc)]
    if suf in (".txt", ".md", ".markdown"):
        text = path.read_text(encoding="utf-8", errors="replace")
        return [("1", text)]
    if suf == ".pptx":
        try:
            from pptx import Presentation      # optional dependency
        except ImportError:
            raise ESError("Reading .pptx needs python-pptx.", "Export the slides as PDF, or: python -m pip install python-pptx")
        prs = Presentation(str(path))
        out = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            out.append((str(i), "\n".join(texts)))
        return out
    raise ESError(f"Unsupported material format {suf}.", "Use PDF, TXT, Markdown or PPTX.")


class _SectionIndex:
    """Tiny TF-IDF over the book's content sections, used to propose candidates."""

    def __init__(self, ws: Workspace):
        self.docs: list[tuple[str, str, Counter]] = []
        for ch in ws.chapters_in_scope():
            for sec in ch["sections"]:
                if sec["kind"] != "content":
                    continue
                text = " ".join(p["text"] for _, p in section_paragraphs(ws, sec))
                self.docs.append((sec["id"], f"{sec['id']} {sec['title']}", Counter(_tokens(text + " " + sec["title"] * 3))))
        self.df: Counter = Counter()
        for _, _, tf in self.docs:
            self.df.update(set(tf))
        self.n = max(1, len(self.docs))
        inv_terms = {}
        for iid, it in ws.all_inventory_items().items():
            if it["kind"] == "term":
                inv_terms[it["text"].casefold()] = it.get("section")
        self.terms = inv_terms

    def rank(self, text: str, k: int = 5) -> list[tuple[str, str, float]]:
        q = Counter(_tokens(text))
        low = text.casefold()
        scores = []
        for sid, title, tf in self.docs:
            s = 0.0
            for w, c in q.items():
                if w in tf:
                    s += (1 + math.log(tf[w])) * math.log(1 + self.n / (1 + self.df[w]))
            for term, tsec in self.terms.items():
                if tsec == sid and len(term) > 3 and term in low:
                    s += 4.0
            scores.append((sid, title, s))
        scores.sort(key=lambda x: -x[2])
        return [x for x in scores[:k] if x[2] > 0] or scores[:3]


def _split_questions(pages: list[tuple[str, str]]) -> list[tuple[str, str, str]]:
    """[(number, page, text)] using common numbering styles."""
    out: list[list[str]] = []
    start = re.compile(r"^\s*(?:Q(?:uestion)?\.?\s*)?(\d{1,3})\s*[.):]\s+(\S.*)$", re.I)
    for label, text in pages:
        for line in text.splitlines():
            m = start.match(line)
            is_new = False
            if m:
                n = int(m.group(1))
                # question numbers must go up by 1-3; this ignores numbered lists inside a question
                is_new = n <= 3 if not out else int(out[-1][0]) < n <= int(out[-1][0]) + 3
            if is_new:
                out.append([m.group(1), label, m.group(2).strip()])
            elif out and line.strip():
                if len(out[-1][2]) < 1500:
                    out[-1][2] += " " + line.strip()
    return [(n, p, re.sub(r"\s+", " ", t).strip()) for n, p, t in out][:80]


def add_material(ws: Workspace, path: Path, kind: str, label: str | None) -> str:
    if not path.exists():
        raise ESError(f"File not found: {path}")
    mats = ws.state.setdefault("materials", {})
    mid = (label or path.stem)
    mid = re.sub(r"[^\w.\-]+", "-", mid).strip("-")[:30] or f"m{len(mats) + 1}"
    if mid in mats:
        mid = f"{mid}-{sha(str(path), 4)}"
    d = ws.materials_dir / mid
    d.mkdir(parents=True, exist_ok=True)
    dest = d / path.name
    shutil.copyfile(path, dest)
    pages = _extract_text(dest)
    if not "".join(t for _, t in pages).strip():
        raise ESError("No text could be extracted from this file (scanned?).",
                      "Provide a text-based PDF or a .txt copy of the questions.")
    rec = {"kind": kind, "label": label or path.stem, "file": str(dest.relative_to(ws.root).as_posix()),
           "status": "worksheet", "task": f"map-material:{mid}"}
    idx = _SectionIndex(ws)
    blocks, originals = [], {}
    if kind == "past-paper":
        qs = _split_questions(pages)
        if not qs:
            raise ESError("No numbered questions found in this paper.",
                          "Numbering like '1.' or 'Question 1' is needed. Or save the questions as a .txt file, one per line, numbered.")
        for n, pg, text in qs:
            bid = f"PQ-{sha(mid, 4)}-{n}"
            cands = idx.rank(text)
            cand_text = " | ".join(f"{sid} {title.split(' ', 1)[1] if ' ' in title else title}" for sid, title, _ in cands)
            originals[bid] = {"question": text[:700], "candidates": cand_text}
            blocks.append("\n".join([f"::: map-question {bid}", f"question: {text[:700]}", f"candidates: {cand_text}",
                                     "pick: <<FILL: 1-3 section numbers from candidates (comma separated), or NONE>>",
                                     "reason: <<FILL: one sentence: which idea the question tests>>", ":::"]))
        header = ("# Past paper mapping — " + rec["label"] + "\n<!-- For each exam question, pick the book "
                  "section(s) it tests, ONLY from its candidates line. NONE if it is not in the book. -->\n")
        kind_ws = "map-question"
    elif kind == "syllabus":
        lines = [l.strip(" \t-*•") for _, t in pages for l in t.splitlines()]
        lines = [l for l in lines if len(l.split()) >= 3][:80]
        for i, line in enumerate(lines, start=1):
            bid = f"SY-{sha(mid, 4)}-{i}"
            cands = idx.rank(line)
            cand_text = " | ".join(f"{sid} {title.split(' ', 1)[1] if ' ' in title else title}" for sid, title, _ in cands)
            originals[bid] = {"syllabus": line[:400], "candidates": cand_text}
            blocks.append("\n".join([f"::: scope-item {bid}", f"syllabus: {line[:400]}", f"candidates: {cand_text}",
                                     "pick: <<FILL: section numbers from candidates that this line covers, or NONE>>",
                                     ":::"]))
        header = ("# Syllabus mapping — " + rec["label"] + "\n<!-- For each syllabus line, pick the book "
                  "sections it covers, ONLY from its candidates line, or NONE. -->\n")
        kind_ws = "scope-item"
    else:
        items = _slide_conflicts(ws, pages)
        if not items:
            rec["status"] = "done"
            rec["note"] = "no key terms from the book appear in these slides"
            ws.state["materials"][mid] = rec
            ws.save_state()
            return f"Added slides '{rec['label']}'. None of the book's key terms appear in them; nothing to compare."
        for tid, term, book, slide in items:
            bid = f"CF-{sha(mid + tid, 6)}"
            originals[bid] = {"term": term, "book": book, "slides": slide}
            blocks.append("\n".join([f"::: conflict {bid}", f"term: {term}", f"book: {book}", f"slides: {slide}",
                                     "verdict: <<FILL: CONSISTENT | DIFFERENT_NOTATION | CONFLICT | NOT_RELATED>>",
                                     "explain: <<FILL: one sentence>>", ":::"]))
            rec.setdefault("terms", {})[bid] = tid
        header = ("# Slides vs book — " + rec["label"] + "\n<!-- CONSISTENT = same meaning. DIFFERENT_NOTATION = "
                  "same idea, different symbols or names. CONFLICT = the slides say something different. NOT_RELATED = "
                  "the slide line is not about this term. -->\n")
        kind_ws = "conflict"
    wpath = d / "worksheet.md"
    write_text(wpath, header + "\n" + "\n\n".join(blocks) + "\n")
    write_json(d / "worksheet.json", {"kind": kind_ws, "originals": originals})
    rec["worksheet"] = str(wpath.relative_to(ws.root).as_posix())
    ws.state["materials"][mid] = rec
    ws.state["plan_confirmed"] = False          # priorities change with new material
    ws.save_state()
    return (f"Added {kind} '{rec['label']}' with {len(originals)} item(s) to map. "
            f"Run next to get the mapping task.")


def _slide_conflicts(ws: Workspace, pages: list[tuple[str, str]]) -> list[tuple[str, str, str, str]]:
    out = []
    items = [it for it in ws.all_inventory_items().values() if it["kind"] == "term"]
    for it in items:
        term = it["text"]
        rx = re.compile(rf"\b{re.escape(term)}\b", re.I)
        best = None
        for label, text in pages:
            for s in split_sentences(text.replace("\n", " ")):
                if rx.search(s):
                    score = 2 if re.search(r"\b(is|are|means|defined|refers to)\b|[:=]", s) else 1
                    if best is None or score > best[0]:
                        best = (score, label, s[:300])
            if best and best[0] == 2:
                break
        if best:
            out.append((it["id"], term, f"[p.{it.get('page')}] {it.get('context', '')[:300]}",
                        f"[slide {best[1]}] {best[2]}"))
    out.sort(key=lambda x: 0 if re.search(r"\b(is|are|means|defined)\b", x[3]) else 1)
    return out[:40]


def pending_material_task(ws: Workspace):
    from .pipeline import Task
    for mid, rec in ws.state.get("materials", {}).items():
        if rec.get("status") != "worksheet":
            continue
        path = ws.root / rec["worksheet"]
        kind = {"past-paper": "map past-exam questions to book sections",
                "syllabus": "map syllabus lines to book sections",
                "slides": "compare the lecture slides with the book"}[rec["kind"]]
        rules = ["Choose ONLY from the candidates line (or NONE)." if rec["kind"] != "slides" else
                 "Judge only from the two quoted lines.",
                 "Keep read-only lines (question:, syllabus:, candidates:, term:, book:, slides:) unchanged.",
                 "Replace every <<FILL ...>>."]
        return Task(rec["task"], "map-material", "careful reader", f"Material '{rec['label']}': {kind}.", edit=path,
                    rules=rules, extra={"mid": mid})
    return None


def check_material_task(ws: Workspace, task) -> tuple[bool, str]:
    from .lint import lint_worksheet
    from .pipeline import _format_issues
    mid = task.extra["mid"]
    rec = ws.state["materials"][mid]
    path = ws.root / rec["worksheet"]
    meta = read_json(path.with_suffix(".json"), {})
    text = read_text(path)
    fixed, _ = autofix(text)
    if fixed != text:
        write_text(path, fixed)
    doc = parse(fixed, path)
    res = lint_worksheet(doc, meta["originals"], meta["kind"])
    if not res.ok:
        return False, _format_issues(ws, path, res.issues, 1, 1)
    if meta["kind"] in ("map-question", "scope-item"):
        mapping = {}
        for blk in doc.blocks:
            picks = [p for p in re.split(r"[,\s]+", blk.value("pick")) if p and p.upper() != "NONE"]
            mapping[blk.id] = picks
        rec["mapping"] = mapping
        msg = f"{sum(1 for v in mapping.values() if v)} of {len(mapping)} items mapped to the book"
    else:
        conflicts = {}
        for blk in doc.blocks:
            v = blk.value("verdict").strip().upper()
            if v in ("CONFLICT", "DIFFERENT_NOTATION"):
                tid = rec["terms"][blk.id]
                conflicts[tid] = {"verdict": v, "explain": blk.value("explain"), "slides": meta["originals"][blk.id]["slides"],
                                  "material": rec["label"]}
        rec["conflicts"] = conflicts
        msg = f"{len(conflicts)} difference(s) between the slides and the book recorded"
    rec["status"] = "done"
    ws.state["plan_confirmed"] = False
    ws.save_state()
    return True, msg


def conflicts_by_term(ws: Workspace) -> dict[str, list[dict]]:
    out: dict[str, list[dict]] = {}
    for rec in ws.state.get("materials", {}).values():
        for tid, c in (rec.get("conflicts") or {}).items():
            out.setdefault(tid, []).append(c)
    return out
